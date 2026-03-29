"""Unit tests for app.export (.docx content export)."""

import io

import pytest
from docx import Document

from app.export import export_docx_bytes
from app.models import (
    AlignmentEnum,
    FontConfig,
    HeaderFooterConfig,
    LayoutConfig,
    MarginConfig,
    PageConfig,
    ParagraphConfig,
    StyleConfig,
)


def _read_para_texts(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs]


def _default_config() -> LayoutConfig:
    return LayoutConfig(
        title="Test Template",
        styles=[
            StyleConfig(name="Normal", font=FontConfig(size=12)),
            StyleConfig(name="Title", font=FontConfig(size=16, bold=True)),
            StyleConfig(name="Heading 1", font=FontConfig(size=14, bold=True)),
            StyleConfig(name="Heading 2", font=FontConfig(size=12, bold=True, italic=True)),
            StyleConfig(name="Heading 3", font=FontConfig(size=12, italic=True)),
            StyleConfig(name="Abstract", font=FontConfig(size=11)),
            StyleConfig(name="References", font=FontConfig(size=10)),
        ],
    )


# ---------------------------------------------------------------------------
# Basic output
# ---------------------------------------------------------------------------


class TestExportDocxBytes:
    def test_returns_bytes(self):
        result = export_docx_bytes("Hello world", _default_config())
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_valid_docx(self):
        result = export_docx_bytes("Hello world", _default_config())
        doc = Document(io.BytesIO(result))
        assert doc is not None

    def test_empty_content(self):
        """Empty content should not raise an error."""
        result = export_docx_bytes("", _default_config())
        assert isinstance(result, bytes)


# ---------------------------------------------------------------------------
# Markup parsing
# ---------------------------------------------------------------------------


class TestMarkupParsing:
    def test_title_heading(self):
        content = "# My Paper Title"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "My Paper Title" in texts

    def test_h2_heading(self):
        content = "## 1. Introduction"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "1. Introduction" in texts

    def test_h3_heading(self):
        content = "### 1.1 Background"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "1.1 Background" in texts

    def test_h4_heading(self):
        content = "#### 1.1.1 Detail"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "1.1.1 Detail" in texts

    def test_body_text(self):
        content = "This is a normal paragraph of text."
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "This is a normal paragraph of text." in texts

    def test_blank_lines_skipped(self):
        content = "Para one\n\nPara two"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "Para one" in texts
        assert "Para two" in texts

    def test_separator_skipped(self):
        content = "Para one\n---\nPara two"
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "---" not in texts

    def test_abstract_section(self):
        content = "Abstract\nThis is the abstract text."
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "Abstract" in texts
        assert "This is the abstract text." in texts

    def test_references_section(self):
        content = "References\n[1] Smith, J. (2023). Test."
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "References" in texts
        assert "[1] Smith, J. (2023). Test." in texts

    def test_numbered_reference_item(self):
        content = "References\n1. Smith, J. (2023). Test."
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "1. Smith, J. (2023). Test." in texts

    def test_full_paper(self):
        content = """\
# Deep Learning in Drug Discovery

Author A, Author B

Abstract
This paper surveys deep learning applications in drug discovery.

## 1. Introduction
Drug discovery is a long and expensive process.

### 1.1 Motivation
Deep learning offers promising acceleration.

## 2. Methods
We survey 50 recent publications.

## References
[1] LeCun, Y. et al. (2015). Deep learning. Nature.
[2] Jumper, J. et al. (2021). AlphaFold2. Nature.
"""
        result = export_docx_bytes(content, _default_config())
        texts = _read_para_texts(result)
        assert "Deep Learning in Drug Discovery" in texts
        assert "1. Introduction" in texts
        assert "1.1 Motivation" in texts
        assert "Drug discovery is a long and expensive process." in texts
        assert "[1] LeCun, Y. et al. (2015). Deep learning. Nature." in texts


# ---------------------------------------------------------------------------
# Config options applied
# ---------------------------------------------------------------------------


class TestConfigApplication:
    def test_two_column_layout_does_not_raise(self):
        config = LayoutConfig(
            page=PageConfig(columns=2, column_spacing=0.5),
            styles=[StyleConfig(name="Normal")],
        )
        result = export_docx_bytes("Some body text here.", config)
        assert len(result) > 0

    def test_header_footer_applied(self):
        config = LayoutConfig(
            header=HeaderFooterConfig(text="Test Journal", include_page_number=True),
            footer=HeaderFooterConfig(include_page_number=True),
        )
        result = export_docx_bytes("Body text.", config)
        assert len(result) > 0

    def test_custom_page_margins(self):
        config = LayoutConfig(
            page=PageConfig(
                margins=MarginConfig(top=2.0, bottom=2.0, left=2.0, right=2.0)
            )
        )
        result = export_docx_bytes("Text.", config)
        assert len(result) > 0
